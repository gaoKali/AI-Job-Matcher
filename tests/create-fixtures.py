"""Fictional local fixtures. Never reads a user's resume."""
from pathlib import Path
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.lib.pdfencrypt import StandardEncryption
from PIL import Image, ImageDraw, ImageFont
from docx import Document

root = Path(__file__).parent / 'fixtures'
root.mkdir(exist_ok=True)
pdfmetrics.registerFont(TTFont('DemoChinese', r'C:\Windows\Fonts\simhei.ttf'))
pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
lines = ['虚构测试候选人 - Demo Candidate', '工作经验：4年 / Experience: 4 years', '核心技能：用户研究、数据分析、项目协作', '2022-2026 示例公司（虚构） - 产品运营', '职责：整理用户反馈，参与需求分析。', '教育背景：示例大学（虚构），本科。']

def pdf(name, content=lines, font='DemoChinese', encryption=None):
    c = canvas.Canvas(str(root / name), encrypt=encryption)
    c.setTitle('Fictional resume fixture')
    c.setFont(font, 14)
    for i, line in enumerate(content):
        c.drawString(45, 780 - i * 34, line)
    c.showPage()
    c.save()

pdf('normal.pdf')
pdf('normal-cid.pdf', font='STSong-Light')
pdf('normal-en.pdf', content=['Demo Candidate', 'Experience: 4 years', 'Skills: Research, Python, SQL', 'Education: Demo University'], font='Helvetica')
pdf('empty.pdf', content=[])
pdf('encrypted.pdf', encryption=StandardEncryption('demo-password'))
c = canvas.Canvas(str(root / 'multi-page.pdf'))
for i in range(2):
    c.setFont('DemoChinese', 14)
    c.drawString(45, 770, f'第{i + 1}页 - 虚构测试简历')
    c.drawString(45, 735, f'经历{i + 1}：项目协作与用户研究')
    c.showPage()
c.save()
c = canvas.Canvas(str(root / 'two-columns.pdf'))
c.setFont('DemoChinese', 12)
c.drawString(45, 790, '虚构双栏测试简历')
for i, (left, right) in enumerate([('核心技能', '工作经历'), ('用户研究', '示例公司'), ('数据分析', '产品运营'), ('教育背景', '项目经历'), ('示例大学', '用户反馈整理')]):
    c.drawString(45, 740-i*30, left)
    c.drawString(310, 740-i*30, right)
c.showPage(); c.save()
image = Image.new('RGB', (800, 1000), 'white')
draw = ImageDraw.Draw(image)
font = ImageFont.truetype(r'C:\Windows\Fonts\simhei.ttf', 24)
for i, line in enumerate(lines):
    draw.text((40, 60+i*70), line, fill='black', font=font)
image.save(root / 'scan.pdf', 'PDF', resolution=100)

doc = Document()
doc.add_heading(lines[0], 0)
for line in lines[1:]:
    doc.add_paragraph(line)
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = '项目经历'; table.cell(0, 1).text = '负责角色'
table.cell(1, 0).text = '示例项目'; table.cell(1, 1).text = '需求整理'
doc.save(root / 'normal.docx')
Document().save(root / 'empty.docx')
(root / 'corrupt.pdf').write_bytes(b'%PDF-1.7\nBroken fictional test file')
(root / 'corrupt.docx').write_bytes(b'PK\x03\x04Broken fictional test file')
(root / 'wrong.pdf').write_text('This is text, not a PDF.', encoding='utf8')
(root / 'unsupported.txt').write_text('Fictional test only', encoding='utf8')
for extension in ['pdf', 'docx']:
    (root / f'zero-byte.{extension}').write_bytes(b'')
with (root / 'oversize.pdf').open('wb') as oversized:
    oversized.truncate(10 * 1024 * 1024 + 1)
print('Created fictional fixtures in', root)
